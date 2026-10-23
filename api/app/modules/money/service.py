"""B14e Money: turn the business's own exports (bank statement, payment gateway, store/POS orders) into transactions,
then compute revenue, operating margin, burn and runway. Benchmarks: Damodaran, Margins by industry (India), Jan 2026."""
import csv
import io
import re
from collections import defaultdict
from datetime import date, datetime
from functools import cache
from pathlib import Path

BENCHMARKS = Path(__file__).resolve().parents[2] / "data" / "damodaran-india-margins-2026.csv"
MAX_ROWS = 20000

# Header words → field. Checked in order; the first matching column wins.
HINTS = {
    "date": ("invoicedate", "transactiondate", "txndate", "valuedate", "orderdate", "createdat", "date", "created"),
    "credit": ("deposit", "credit", "moneyin", "received"),
    "debit": ("withdrawal", "debit", "moneyout", "paid"),
    "quantity": ("quantity", "qty", "units"),
    "unit_price": ("unitprice", "price", "rate", "mrp"),
    "amount": ("amount", "total", "netamount", "grandtotal", "value", "sales"),
    "product": ("product", "item", "description", "sku", "lineitem"),
    "description": ("narration", "particulars", "description", "remarks", "details", "note"),
    "kind": ("type", "category", "kind"),
}


def parse_amount(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    neg = s.startswith("(") and s.endswith(")") or s.startswith("-")
    s = re.sub(r"[^\d.]", "", s.replace("Rs.", "").replace("Rs", ""))
    if not s or s == ".":
        return None
    return -float(s) if neg else float(s)


def parse_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    if not s:
        return None
    if re.match(r"\d{4}-\d{2}-\d{2}", s):
        return date.fromisoformat(s[:10])
    m = re.match(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", s)  # day first, as Indian banks and stores write it
    if m:
        d, mo, y = (int(x) for x in m.groups())
        return date(y + 2000 if y < 100 else y, mo, d)
    return None


def parse_table(filename: str, data: bytes) -> tuple[list[str], list[dict]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else f"Column {i + 1}" for i, h in enumerate(next(it))]
        rows = [dict(zip(headers, r)) for r, _ in zip(it, range(MAX_ROWS)) if any(c is not None for c in r)]
        return headers, rows
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    headers = [h.strip() for h in reader.fieldnames or []]
    rows = [{k.strip(): v for k, v in r.items() if k} for r, _ in zip(reader, range(MAX_ROWS))]
    return headers, rows


def suggest_mapping(headers: list[str]) -> dict:
    norm = {h: re.sub(r"[^a-z]", "", h.lower()) for h in headers}
    out: dict[str, str] = {}
    for field, words in HINTS.items():
        for w in words:
            col = next((h for h in headers if w in norm[h] and (h not in out.values() or field in ("description", "product"))), None)
            if col:
                out[field] = col
                break
    if "credit" in out or "debit" in out:  # bank statement: the two columns are the amounts
        out.pop("amount", None), out.pop("quantity", None), out.pop("unit_price", None), out.pop("product", None)
    elif "quantity" in out and "unit_price" in out:  # order lines: amount = qty × price
        out.pop("amount", None)
    else:
        out.pop("unit_price", None)
    if out.get("product") == out.get("description") and "quantity" not in out:
        out.pop("product")  # free-text narration isn't a product name
    return out


def _kind(raw: str | None, amount: float) -> str:
    s = (raw or "").lower()
    if any(w in s for w in ("fund", "investment", "equity", "grant")):
        return "funding"
    if any(w in s for w in ("expense", "purchase", "payment", "debit", "cost", "salary", "rent")):
        return "expense"
    return "sale"


def to_transactions(rows: list[dict], mapping: dict, source: str) -> list[dict]:
    out = []
    g = lambda r, f: r.get(mapping[f]) if f in mapping else None
    for r in rows:
        when = parse_date(g(r, "date"))
        if not when:
            continue
        product = str(g(r, "product") or "").strip() or None
        desc = str(g(r, "description") or "").strip() or None
        qty = parse_amount(g(r, "quantity"))
        if "credit" in mapping or "debit" in mapping:
            credit, debit = parse_amount(g(r, "credit")), parse_amount(g(r, "debit"))
            if credit:
                out.append({"date": when, "kind": _kind(desc, credit) if "fund" in (desc or "").lower() else "sale", "amount": credit,
                            "product": None, "quantity": None, "description": desc, "source": source})
            if debit:
                out.append({"date": when, "kind": "expense", "amount": debit, "product": None, "quantity": None, "description": desc, "source": source})
            continue
        if "unit_price" in mapping:
            price = parse_amount(g(r, "unit_price"))
            amount = round(qty * price, 2) if qty is not None and price is not None else None
        else:
            amount = parse_amount(g(r, "amount"))
        if amount is None:
            continue
        kind = _kind(str(g(r, "kind") or ""), amount) if "kind" in mapping else "sale"
        out.append({"date": when, "kind": kind, "amount": abs(amount) if kind != "sale" else amount, "product": product,
                    "quantity": qty, "description": desc, "source": source})
    return out


def metrics(txs: list[dict], cash: float | None) -> dict:
    by_month: dict[str, dict] = defaultdict(lambda: {"revenue": 0.0, "expenses": 0.0})
    funding, products = 0.0, defaultdict(lambda: {"units": 0.0, "revenue": 0.0})
    for t in txs:
        if t["kind"] == "funding":
            funding += t["amount"]
            continue
        key = t["date"].strftime("%Y-%m")
        by_month[key]["revenue" if t["kind"] == "sale" else "expenses"] += t["amount"]
        if t["kind"] == "sale" and t.get("product"):
            products[t["product"]]["units"] += t.get("quantity") or 0
            products[t["product"]]["revenue"] += t["amount"]
    months = [{"month": k, "revenue": round(v["revenue"]), "expenses": round(v["expenses"])} for k, v in sorted(by_month.items())]
    last = months[-3:]
    rev = sum(m["revenue"] for m in last) / len(last) if last else 0
    exp = sum(m["expenses"] for m in last) / len(last) if last else 0
    burn = max(exp - rev, 0)
    return {
        "months": months,
        "monthly_revenue": round(rev), "monthly_expenses": round(exp),
        "operating_margin": round((rev - exp) / rev, 4) if rev else None,
        "monthly_burn": round(burn),
        "runway_months": round(cash / burn, 1) if burn and cash is not None else None,
        "funding_total": round(funding),
        "products": sorted(({"product": k, "units": round(v["units"]), "revenue": round(v["revenue"]),
                             "avg_price": round(v["revenue"] / v["units"]) if v["units"] else None} for k, v in products.items()),
                           key=lambda p: -p["revenue"]),
    }


@cache
def _benchmarks() -> dict[str, dict]:
    with BENCHMARKS.open(encoding="utf-8") as f:
        return {r["industry"]: {"industry": r["industry"], "firms": int(r["firms"]),
                                **{k: float(r[k]) if r[k] else None for k in ("gross_margin", "net_margin", "operating_margin", "ebitda_margin")}}
                for r in csv.DictReader(f)}


def benchmark(industry: str | None) -> dict | None:
    return _benchmarks().get(industry or "")


def industries() -> list[str]:
    return sorted(_benchmarks())
